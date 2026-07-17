from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx2pdf import convert
import os

doc = Document()
sec1 = doc.sections[0]
doc.add_paragraph('Section 1 - Title Page')

xml_str = f'<w:pgBorders {nsdecls("w")} w:offsetFrom="page"><w:top w:val="single" w:sz="8" w:space="24" w:color="000000"/><w:left w:val="single" w:sz="8" w:space="24" w:color="000000"/><w:bottom w:val="single" w:sz="8" w:space="24" w:color="000000"/><w:right w:val="single" w:sz="8" w:space="24" w:color="000000"/></w:pgBorders>'

sec1._sectPr.append(parse_xml(xml_str))

sec2 = doc.add_section()
doc.add_paragraph('Section 2 - Certificate Page')
sec2._sectPr.append(parse_xml(xml_str))

sec3 = doc.add_section()
doc.add_paragraph('Section 3 - Abstract Page (No Border)')

doc.save('test_border.docx')
convert('test_border.docx', 'test_border.pdf')
print('SUCCESS: Generated test_border.pdf')
